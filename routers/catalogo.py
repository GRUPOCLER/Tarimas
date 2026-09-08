from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Header
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
from typing import Optional
from urllib.parse import quote
import io, os, csv, re, httpx

from database import get_db
from models.models import CatalogoItem
from services.auth import verificar_token

router = APIRouter()

# Hoja de Google Sheets con el catalogo — via variables de entorno para
# poder cambiarla sin tocar codigo. Requiere que este compartida como
# "Cualquier persona con el enlace" (solo lectura), sin eso Google
# regresa una pagina de login en vez del CSV.
GOOGLE_SHEET_CATALOGO_ID   = os.getenv("GOOGLE_SHEET_CATALOGO_ID", "12Ewr2SvJXRt-g5HqGo-JrnJ07mv3vXboHU7Y2Gfkedw")
GOOGLE_SHEET_CATALOGO_HOJA = os.getenv("GOOGLE_SHEET_CATALOGO_HOJA", "CATALOGO_V2")

async def get_current_user(authorization: str = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Sin autorizacion")
    payload = verificar_token(authorization.split(" ")[1])
    if not payload:
        raise HTTPException(status_code=401, detail="Token invalido")
    return payload

# Nombres de columna aceptados en el Excel/CSV — flexible por si el
# archivo real trae variantes de mayusculas/espacios/acentos.
_ALIAS = {
    "sku": "sku", "clave": "sku", "codigo de articulo": "sku",
    "descripcion": "descripcion", "descripción": "descripcion",
    "peso_unit": "peso_unit", "peso unit": "peso_unit", "peso unitario": "peso_unit", "peso": "peso_unit",
    "lr_unit": "lr_unit", "lr unit": "lr_unit", "largo": "lr_unit", "largo unitario": "lr_unit",
    "an_unit": "an_unit", "an unit": "an_unit", "ancho": "an_unit", "ancho unitario": "an_unit",
    "al_unit": "al_unit", "al unit": "al_unit", "alto": "al_unit", "alto unitario": "al_unit",
    "m3_unit": "m3_unit", "m3 unit": "m3_unit", "m3": "m3_unit",
    "cm_cant": "cm_cant", "cm cant": "cm_cant", "cantidad caja master": "cm_cant", "pzas por caja": "cm_cant", "piezas por caja": "cm_cant",
    "cm_peso": "cm_peso", "cm peso": "cm_peso", "peso caja master": "cm_peso",
    "cm_lr": "cm_lr", "cm lr": "cm_lr", "largo caja master": "cm_lr",
    "cm_an": "cm_an", "cm an": "cm_an", "ancho caja master": "cm_an",
    "cm_al": "cm_al", "cm al": "cm_al", "alto caja master": "cm_al",
    "cm_m3": "cm_m3", "cm m3": "cm_m3", "m3 caja master": "cm_m3",
    "sat_codigo": "sat_codigo", "sat codigo": "sat_codigo", "codigo sat": "sat_codigo", "código sat": "sat_codigo",
    "sat_desc": "sat_desc", "sat descripcion": "sat_desc", "descripcion sat": "sat_desc",
}

def _normalizar_col(nombre: str) -> str:
    n = str(nombre).strip().lower()
    n = re.sub(r"\([^)]*\)", "", n).strip()  # quita "(kg)", "(cm)", etc.
    n = re.sub(r"\s+", " ", n)
    return _ALIAS.get(n, n.replace(" ", "_"))

def _num(v, es_entero=False):
    if v is None or v == "":
        return 0
    try:
        n = float(str(v).replace(",", "").strip())
        return int(round(n)) if es_entero else n
    except (ValueError, TypeError):
        return 0

async def _upsert_catalogo(db: AsyncSession, filas: list) -> dict:
    """Recibe una lista de dicts ya normalizados (claves como 'sku', 'cm_cant', etc.)
    y hace insert/update en la tabla catalogo. Compartido entre /importar y
    /sincronizar-drive para no duplicar la logica."""
    insertados, actualizados, omitidos = 0, 0, 0
    for fila in filas:
        sku = str(fila.get("sku") or "").strip().upper()
        if not sku:
            omitidos += 1
            continue

        existente = await db.execute(select(CatalogoItem).where(CatalogoItem.sku == sku))
        item = existente.scalar_one_or_none()
        es_nuevo = item is None
        if es_nuevo:
            item = CatalogoItem(sku=sku)
            db.add(item)

        item.descripcion = str(fila.get("descripcion") or (item.descripcion or ""))
        item.peso_unit   = _num(fila.get("peso_unit"))
        item.lr_unit     = _num(fila.get("lr_unit"))
        item.an_unit     = _num(fila.get("an_unit"))
        item.al_unit     = _num(fila.get("al_unit"))
        item.m3_unit     = _num(fila.get("m3_unit"))
        item.cm_cant     = _num(fila.get("cm_cant"), es_entero=True)
        item.cm_peso     = _num(fila.get("cm_peso"))
        item.cm_lr       = _num(fila.get("cm_lr"))
        item.cm_an       = _num(fila.get("cm_an"))
        item.cm_al       = _num(fila.get("cm_al"))
        item.cm_m3       = _num(fila.get("cm_m3"))
        item.sat_codigo  = str(fila.get("sat_codigo") or "") or None
        item.sat_desc    = str(fila.get("sat_desc") or "") or None

        if es_nuevo: insertados += 1
        else: actualizados += 1

    await db.commit()
    return {"insertados": insertados, "actualizados": actualizados, "omitidos": omitidos, "total_filas": len(filas)}

@router.get("/")
async def listar_catalogo(
    buscar: Optional[str] = None,
    limite: int = 100,
    db: AsyncSession = Depends(get_db),
    user: dict = Depends(get_current_user)
):
    q = select(CatalogoItem).order_by(CatalogoItem.sku).limit(limite)
    if buscar:
        patron = f"%{buscar.strip()}%"
        q = q.where(CatalogoItem.sku.ilike(patron) | CatalogoItem.descripcion.ilike(patron))
    result = await db.execute(q)
    items = result.scalars().all()
    return [{
        "sku": i.sku, "descripcion": i.descripcion,
        "peso_unit": i.peso_unit, "lr_unit": i.lr_unit, "an_unit": i.an_unit, "al_unit": i.al_unit,
        "cm_cant": i.cm_cant, "cm_peso": i.cm_peso, "cm_lr": i.cm_lr, "cm_an": i.cm_an, "cm_al": i.cm_al,
        "sat_codigo": i.sat_codigo,
    } for i in items]

@router.get("/contar")
async def contar_catalogo(db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    result = await db.execute(select(func.count(CatalogoItem.sku)))
    con_master = await db.execute(select(func.count(CatalogoItem.sku)).where(CatalogoItem.cm_cant > 0))
    return {"total": result.scalar() or 0, "con_caja_master": con_master.scalar() or 0}

@router.post("/importar")
async def importar_catalogo(
    archivo: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: dict = Depends(get_current_user)
):
    if user.get("rol") != "admin":
        raise HTTPException(status_code=403, detail="Solo un Administrador puede importar el catalogo")

    contenido = await archivo.read()
    nombre = (archivo.filename or "").lower()

    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="Falta la libreria openpyxl en el servidor")

    filas = []
    hoja_usada = "(csv)"
    if nombre.endswith(".csv"):
        texto = contenido.decode("utf-8-sig", errors="ignore")
        lector = csv.reader(io.StringIO(texto))
        encabezado = None
        for fila in lector:
            if encabezado is None:
                encabezado = [_normalizar_col(c) for c in fila]
                continue
            filas.append(dict(zip(encabezado, fila)))
    else:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
        ws = None
        for nombre_hoja in wb.sheetnames:
            if "catalogo" in nombre_hoja.strip().lower():
                ws = wb[nombre_hoja]
                break
        if ws is None:
            ws = wb.active  # respaldo: si no encuentra una hoja "catalogo", usa la activa
        hoja_usada = ws.title
        filas_raw = list(ws.iter_rows(values_only=True))
        if not filas_raw:
            raise HTTPException(status_code=422, detail="El archivo esta vacio")
        encabezado = [_normalizar_col(c) for c in filas_raw[0]]
        for fila in filas_raw[1:]:
            filas.append(dict(zip(encabezado, fila)))

    resultado = await _upsert_catalogo(db, filas)
    resultado["ok"] = True
    resultado["hoja_usada"] = hoja_usada
    return resultado

@router.post("/sincronizar-drive")
async def sincronizar_desde_drive(db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    """Lee el catalogo directamente desde el Google Sheet compartido — sin
    necesidad de subir un archivo. Requiere que la hoja este compartida como
    'Cualquier persona con el enlace' (solo lectura)."""
    if user.get("rol") != "admin":
        raise HTTPException(status_code=403, detail="Solo un Administrador puede sincronizar el catalogo")
    if not GOOGLE_SHEET_CATALOGO_ID:
        raise HTTPException(status_code=500, detail="GOOGLE_SHEET_CATALOGO_ID no esta configurado en Railway")

    url = (
        f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_CATALOGO_ID}/gviz/tq"
        f"?tqx=out:csv&sheet={quote(GOOGLE_SHEET_CATALOGO_HOJA)}"
    )
    async with httpx.AsyncClient(timeout=30) as client:
        try:
            r = await client.get(url, follow_redirects=True)
        except Exception as e:
            raise HTTPException(status_code=502, detail=f"No se pudo conectar a Google Sheets: {e}")

    if r.status_code != 200:
        raise HTTPException(
            status_code=502,
            detail=f"Google Sheets respondio {r.status_code} — revisa que la hoja siga compartida "
                   f"como 'Cualquier persona con el enlace' y que la pestana se llame "
                   f"'{GOOGLE_SHEET_CATALOGO_HOJA}'"
        )
    if "<html" in r.text[:200].lower():
        raise HTTPException(
            status_code=502,
            detail="Google regreso una pagina de login en vez del CSV — la hoja dejo de ser publica, "
                   "revisa el permiso de 'Compartir'"
        )

    texto = r.content.decode("utf-8-sig", errors="ignore")
    lector = csv.reader(io.StringIO(texto))
    filas_raw = list(lector)
    if not filas_raw:
        raise HTTPException(status_code=422, detail=f"La pestana '{GOOGLE_SHEET_CATALOGO_HOJA}' esta vacia")

    encabezado = [_normalizar_col(c) for c in filas_raw[0]]
    filas = [dict(zip(encabezado, f)) for f in filas_raw[1:]]

    resultado = await _upsert_catalogo(db, filas)
    resultado["ok"] = True
    resultado["fuente"] = "google_sheets"
    resultado["hoja_usada"] = GOOGLE_SHEET_CATALOGO_HOJA
    return resultado
